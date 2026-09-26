"""Generate the Word solution-overview document for the FE BAR submission."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

DB_RED = RGBColor(0xC4, 0x23, 0x1A)
NAVY = RGBColor(0x1F, 0x33, 0x5B)

doc = Document()
styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)


def h(text, level=1, color=NAVY):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = color
    return p


def para(text, bold=False, italic=False, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(hd)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t


# ---------------------------------------------------------------- title -----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Medical Omnichannel Intelligence — Daiichi Sankyo")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = DB_RED
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("Databricks FE BAR build — solution overview & asset guide")
rs.italic = True
rs.font.size = Pt(12)
para("All data is synthetic. Context is Medical Affairs (non-promotional). Every AI output "
     "is a recommendation for human (MSL) review. Built with Claude (Isaac) as an AI pair-builder.",
     italic=True, size=9)

# ---------------------------------------------------------------- summary ---
h("1. Executive summary", 1)
para("Daiichi Sankyo's oncology franchise (Enhertu / trastuzumab deruxtecan, Datroway / "
     "datopotamab deruxtecan) is advanced through scientific exchange between Medical Science "
     "Liaisons (MSLs) and oncology KOLs. Those engagement signals are siloed across MSL field "
     "notes, medical inquiries, congress activity and digital channels, so MSLs miss follow-ups, "
     "respond slowly to medical inquiries, and under-serve high-influence KOLs.")
para("This build unifies every signal into one governed Databricks platform and turns it into a "
     "prioritized, AI-assisted next-best-engagement workflow surfaced in an MSL cockpit — an "
     "integrated journey across Lakeflow, Unity Catalog, Lakebase, Gen AI, Genie and a Databricks App.")
para("Estimated value: ~$3.4M / year for a 120-MSL field-medical org.", bold=True)
table(["Value lever", "Basis", "Est. annual value"],
      [["MSL prep time saved", "AI briefings replace ~2 h/MSL/week", "~$1.3M"],
       ["Faster medical-inquiry resolution", "Overdue surfacing + SLA analytics", "~$0.6M"],
       ["Priority-KOL reach uplift", "90-day reach ~44% → 60%", "~$1.5M"]])

# ---------------------------------------------------------------- access ----
h("2. Environment & how to access everything", 1)
table(["Item", "Value"],
      [["Workspace", "https://adb-984752964297111.11.azuredatabricks.net (Azure)"],
       ["Unity Catalog schema", "_sa701.moa"],
       ["Raw data Volume", "/Volumes/_sa701/moa/raw/ (9 dataset subfolders)"],
       ["Lakeflow pipeline", "moa_medallion_febar (id 8361155d-f4bb-4a33-b1c1-3219f04fef33)"],
       ["SQL warehouse", "148ccb90800933a1"],
       ["LLM endpoint", "databricks-meta-llama-3-3-70b-instruct"],
       ["Lakebase instance", "moa-medical-copilot (PG16, CU_1)"],
       ["Lakebase UC catalog", "moa_lakebase (database databricks_postgres)"],
       ["Genie space", "Medical Omnichannel Intelligence — Daiichi Sankyo (01f1b9f40b3d1499a1dde849fe4ecaf0)"],
       ["Databricks App", "moa-medical-copilot"],
       ["App URL", "https://moa-medical-copilot-984752964297111.11.azure.databricksapps.com"],
       ["AI/BI dashboard", "Medical Omnichannel Intelligence — Daiichi Sankyo (01f1b9f6cf2119cdb68bbd5df75d577e)"],
       ["GitHub repo", "https://github.com/srihari-db/Medical-OmniChannel-Analytics-FE"]])

# ---------------------------------------------------------------- assets ----
h("3. Asset inventory (what was created and where)", 1)
h("3.1 Unity Catalog data layer (_sa701.moa)", 2)
bullet("Bronze (9 streaming tables): bronze_hcp_master, bronze_msl_interactions, "
       "bronze_medical_inquiries, bronze_congress_events, bronze_digital_engagement, "
       "bronze_publications_trials, bronze_patient_claims, bronze_approved_content, bronze_content_shared")
bullet("Silver (9 materialized views): silver_hcp, silver_medical_inquiries, silver_congress, "
       "silver_digital, silver_publications, silver_claims, silver_approved_content, "
       "silver_content_shared, and silver_msl_extracted (Gen AI extraction from raw notes)")
bullet("Gold (13 materialized views): gold_hcp_360, gold_hcp_priority, gold_next_best_engagement, "
       "gold_medical_topic_interest, gold_msl_interaction_summary, gold_hcp_channel_engagement, "
       "gold_congress_engagement, gold_medical_information_requests, gold_territory_opportunity, "
       "gold_approved_content, gold_content_shared, gold_publications_trials, + serving_msl_action_queue")
bullet("Metric views: metric_hcp_engagement, metric_medical_inquiries (governed KPIs)")
bullet("Volumes: raw (source CSVs), source_documents")
bullet("Governance: schema tags (domain=clinical, data_classification=internal); "
       "gold_hcp_360.npi classified PII; layer/serving tags on operational gold tables; "
       "least-privilege grants to the app service principal")

h("3.2 Lakebase (operational serving)", 2)
bullet("public.msl_action_queue — reverse-ETL synced table from _sa701.moa.serving_msl_action_queue (300 rows)")
bullet("public.msl_review_actions — transactional MSL review write-back table")

h("3.3 Gen AI", 2)
bullet("silver_msl_extracted built by ai_query (Llama 3.3 70B) inside the pipeline — "
       "10+ structured fields + topic_category from each free-text MSL note")
bullet("Pre-engagement briefings generated on demand by the app via ai_query")

h("3.4 Genie & App", 2)
bullet("Genie space over 7 gold tables + 2 metric views, with curated sample questions")
bullet("Medical Engagement Copilot app (app/app.py, backend.py, app.yaml, requirements.txt): "
       "MSL Action Queue tab (Lakebase) + HCP Pre-Engagement tab (360 + AI briefing + write-back)")
bullet("AI/BI dashboard 'Medical Omnichannel Intelligence — Daiichi Sankyo' (dashboard/"
       "medical_omnichannel_intelligence.lvdash.json, published): 6 KPI counters, reach-by-territory, "
       "priority-tier pie, topic-interest bar, overdue-by-product bar, next-best-engagement table")

# ---------------------------------------------------------------- repo ------
h("4. Repository map (GitHub)", 1)
table(["Path", "Contents"],
      [["data_generation/", "generate_synthetic_data.py — seeded Daiichi synthetic data generator"],
       ["data/raw/", "the 9 generated CSV datasets"],
       ["pipeline/transformations/", "01_bronze.sql, 02_silver.sql, 03_gold.sql (Lakeflow medallion)"],
       ["sql/", "00_schema_and_volume.sql, 01_metric_views.sql, 02_ai_query_extraction.sql"],
       ["lakebase/", "setup_lakebase.py, grant_app_access.py"],
       ["genie/", "medical_omnichannel_intelligence.genie.json"],
       ["app/", "Databricks App (Streamlit) — Medical Engagement Copilot"],
       ["evidence/", "01–09 execution evidence (readable as text)"],
       ["docs/", "DECISIONS_AND_TRADEOFFS.md, this Word document"],
       ["deck/", "business presentation (Markdown + PDF)"]])

# ---------------------------------------------------------------- evidence --
h("5. Evidence the build actually ran", 1)
bullet("01 Lakeflow pipeline COMPLETED in 214.66s, 0 errors")
bullet("02 Row counts: 3,081 synthetic rows → bronze/silver/13 gold tables")
bullet("03 Gen AI extraction: 500 notes → 500 structured rows (before/after samples)")
bullet("04 Live grounded AI briefing for the top-priority KOL")
bullet("05 Governed metric view results by territory + next-best-engagement top priorities")
bullet("06 Lakebase: synced queue (300 rows) read + write-back rows + SP grants")
bullet("07 Genie: live NL question → correct SQL → NL answer")
bullet("08 Unity Catalog governance: enforced tag policies, PII classification, scoped grants")
bullet("09 App deployed and ComputeState.ACTIVE, deployment SUCCEEDED")

# ---------------------------------------------------------------- repro -----
h("6. How to reproduce", 1)
for i, s in enumerate([
    "Generate synthetic data: python data_generation/generate_synthetic_data.py --out data/raw",
    "Create UC Volume _sa701.moa.raw; upload each CSV to its own subfolder",
    "Deploy & full-refresh the serverless Lakeflow pipeline from pipeline/transformations/",
    "Create metric views (sql/01_metric_views.sql)",
    "Provision Lakebase, register catalog moa_lakebase, create the synced action queue, "
    "run lakebase/setup_lakebase.py and grant_app_access.py",
    "Create the Genie space over the gold tables + metric views",
    "Deploy the app (app/) with sql_warehouse + lakebase resources"], 1):
    doc.add_paragraph(f"{i}. {s}", style="List Number")

para("Note on tooling: the workspace enforces UC tag policies (approved values only); metric-view "
     "YAML must not use backticks in source:; use pg8000 (pure-Python) for Lakebase setup on Spark "
     "serverless (psycopg[binary] SIGABRTs there) while the app runtime uses psycopg; pin "
     "databricks-sdk>=0.81.0 for the w.database API.", italic=True, size=9)

out = "/Users/srihari.a/febar/Medical-OmniChannel-Analytics-FE/docs/Medical_Omnichannel_Daiichi_Solution_Overview.docx"
doc.save(out)
print("saved", out)
