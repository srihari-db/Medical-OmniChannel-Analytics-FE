"""Augment the FE BAR deck to satisfy the submission rubric:
  - lead with the business OUTCOME (new slide 2)
  - quantify impact in the buyer's KPIs, baseline -> target (rework "Why it matters")
  - frame value for exec sponsor AND domain owner (outcome slide band)
  - document DECISIONS & TRADE-OFFS (new slide)
  - show AI AS A FORCE MULTIPLIER: tools, prompts, patterns (new slide)
  - fix closing order + page-number badges

New slides are cloned at the XML level from design-matched donor slides so they
inherit the exact template look. All shapes are text/auto-shapes (no images),
so a deep-copy of the shape elements is safe.
"""
import copy
from pptx import Presentation
from pptx.util import Emu

SRC = "/Users/srihari.a/febar/Medical-OmniChannel-Analytics-FE/deck/Medical_Omnichannel_Daiichi_Deck.pptx"

SHAPE_TAGS = {"sp", "pic", "graphicFrame", "grpSp", "cxnSp"}


def clone_slide(prs, index):
    """Clone a slide (same layout) by deep-copying its shape elements."""
    src = prs.slides[index]
    new = prs.slides.add_slide(src.slide_layout)
    st = new.shapes._spTree
    # drop placeholders that add_slide injected from the layout
    for sp in list(new.shapes):
        sp._element.getparent().remove(sp._element)
    # copy every real shape element from the donor
    for el in list(src.shapes._spTree):
        tag = el.tag.split("}")[-1]
        if tag in SHAPE_TAGS:
            st.append(copy.deepcopy(el))
    return new


def set_text(slide, name, new):
    shp = next((s for s in slide.shapes if s.name == name), None)
    if shp is None or not shp.has_text_frame:
        raise KeyError(f"missing shape {name!r} on slide")
    tf = shp.text_frame
    first = tf.paragraphs[0]
    if first.runs:
        first.runs[0].text = new
        for r in first.runs[1:]:
            r.text = ""
    else:
        first.text = new
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""


prs = Presentation(SRC)
assert len(prs.slides) == 12, f"expected 12 slides, got {len(prs.slides)} (run once, on the 12-slide deck)"

FOOTER = ("Medical Omnichannel Intelligence  ·  Daiichi Sankyo (synthetic)  ·  "
          "Databricks  ·  _sa701.moa")

# ---------------------------------------------------------------- 1. OUTCOME (donor = idx 9, "Why it matters")
outcome = clone_slide(prs, 9)
set_text(outcome, "Google Shape;393;p22", "THE BUSINESS OUTCOME")
set_text(outcome, "Google Shape;394;p22",
         "One prepared, governed view of every KOL — ~$3.4M/year of impact")
set_text(outcome, "Google Shape;395;p22", "02")
# card 1
set_text(outcome, "Google Shape;399;p22", "360° KOL view")
set_text(outcome, "Google Shape;400;p22", "unified signal")
set_text(outcome, "Google Shape;401;p22",
         "MSL notes, inquiries, congress, digital, pubs & trials in one governed profile.")
# card 2
set_text(outcome, "Google Shape;404;p22", "Prepared MSLs")
set_text(outcome, "Google Shape;405;p22", "minutes, not hours")
set_text(outcome, "Google Shape;406;p22",
         "AI pre-engagement briefing + next-best action for every KOL visit.")
# card 3
set_text(outcome, "Google Shape;409;p22", "Governed & safe")
set_text(outcome, "Google Shape;410;p22", "audit-ready")
set_text(outcome, "Google Shape;411;p22",
         "Non-promotional, human-in-the-loop, one Unity Catalog security model.")
# card 4
set_text(outcome, "Google Shape;414;p22", "~$3.4M / yr")
set_text(outcome, "Google Shape;415;p22", "estimated value")
set_text(outcome, "Google Shape;416;p22",
         "For a 120-MSL field-medical org — KPI detail on the following slides.")
# value band -> persona framing
set_text(outcome, "Google Shape;418;p22", "For the sponsor and the field")
set_text(outcome, "Google Shape;419;p22",
         "Executive sponsor (VP Medical Affairs): quantified ROI, faster inquiry SLAs, and "
         "audit-ready Unity Catalog governance for a regulated, non-promotional function.   "
         "Domain owner (field-medical lead): every MSL arrives prepared, prioritized, and "
         "armed with approved content.")
set_text(outcome, "Google Shape;420;p22", FOOTER)

# ---------------------------------------------------------------- 2. Rework "Why it matters" (idx 9) -> KPI baseline->target
wm = prs.slides[9]
set_text(wm, "Google Shape;393;p22", "WHY IT MATTERS")
set_text(wm, "Google Shape;394;p22", "Impact in your KPIs — baseline → target")
set_text(wm, "Google Shape;395;p22", "11")
set_text(wm, "Google Shape;399;p22", "Reach")
set_text(wm, "Google Shape;400;p22", "62% → 80%")
set_text(wm, "Google Shape;401;p22",
         "Priority-KOL 90-day reach, lifted by under-engaged flags + priority score. ~$1.5M/yr.")
set_text(wm, "Google Shape;404;p22", "Speed")
set_text(wm, "Google Shape;405;p22", "~2h → mins")
set_text(wm, "Google Shape;406;p22",
         "MSL prep per engagement — AI briefing replaces manual context assembly. ~$1.3M/yr.")
set_text(wm, "Google Shape;409;p22", "Service")
set_text(wm, "Google Shape;410;p22", "71% → 95%")
set_text(wm, "Google Shape;411;p22",
         "Medical-inquiry SLA compliance — overdue items flagged by product. ~$0.6M/yr.")
set_text(wm, "Google Shape;414;p22", "Insight")
set_text(wm, "Google Shape;415;p22", "0 → 500+")
set_text(wm, "Google Shape;416;p22",
         "Free-text MSL notes turned into queryable topic-interest signal.")
set_text(wm, "Google Shape;418;p22", "The ROI, quantified")
set_text(wm, "Google Shape;419;p22",
         "~$3.4M / year estimated for a 120-MSL field-medical org — faster prep, sharper "
         "prioritization and faster SLAs — on one governed Lakehouse, no data copies, one "
         "security model.")

# ---------------------------------------------------------------- 3. DECISIONS & TRADE-OFFS (donor = idx 11, "Trust", 6 cards)
TRUST_CARDS = [  # (icon_shape, title_shape, body_shape)
    ("Google Shape;440;p24", "Google Shape;441;p24", "Google Shape;442;p24"),
    ("Google Shape;444;p24", "Google Shape;445;p24", "Google Shape;446;p24"),
    ("Google Shape;448;p24", "Google Shape;449;p24", "Google Shape;450;p24"),
    ("Google Shape;452;p24", "Google Shape;453;p24", "Google Shape;454;p24"),
    ("Google Shape;456;p24", "Google Shape;457;p24", "Google Shape;458;p24"),
    ("Google Shape;460;p24", "Google Shape;461;p24", "Google Shape;462;p24"),
]

tradeoffs = clone_slide(prs, 11)
set_text(tradeoffs, "Google Shape;435;p24", "DECISIONS & TRADE-OFFS")
set_text(tradeoffs, "Google Shape;436;p24", "Why this approach — and what we consciously traded")
set_text(tradeoffs, "Google Shape;437;p24", "12")
TRADEOFFS = [
    ("🗄️", "Lakebase for serving",
     "The app needs single-row reads + transactional write-back — Delta Gold can't hit that latency. "
     "Trade-off: a small Postgres surface, synced from UC."),
    ("🧠", "ai_query() in-pipeline",
     "Extraction stays in governed SQL, no data egress. An external LLM API adds movement, cost & "
     "compliance risk. Trade-off: hosted models only."),
    ("💬", "Genie AND a dashboard",
     "Genie for open-ended questions, AI/BI for always-on KPIs — same governed metrics. One alone "
     "under-serves either explorers or operators."),
    ("🧬", "Llama 3.3 70B, hosted",
     "In-workspace, no third-party data sharing, strong extraction. Trade-off: not the frontier model — "
     "ample for structured extraction."),
    ("🔒", "One UC schema",
     "A single medallion, one lineage graph & security model — vs. fragmented marts. Trade-off: tighter "
     "coordination on a shared schema."),
    ("🧪", "Synthetic data",
     "Safe, reproducible, no PHI/PII — vs. real data that blocks sharing. Trade-off: value figures are "
     "modeled estimates, not measured."),
]
for (icon_s, title_s, body_s), (icon, title, body) in zip(TRUST_CARDS, TRADEOFFS):
    set_text(tradeoffs, icon_s, icon)
    set_text(tradeoffs, title_s, title)
    set_text(tradeoffs, body_s, body)
set_text(tradeoffs, "Google Shape;463;p24", FOOTER)

# ---------------------------------------------------------------- 4. BUILT WITH AI (donor = idx 11 again)
builtai = clone_slide(prs, 11)
set_text(builtai, "Google Shape;435;p24", "BUILT WITH AI")
set_text(builtai, "Google Shape;436;p24", "AI as a force multiplier — tools, prompts & patterns")
set_text(builtai, "Google Shape;437;p24", "13")
BUILT = [
    ("🤖", "Claude Code / Isaac",
     "An agentic AI teammate drove the build end to end — provisioning, code and this deck — "
     "not just autocomplete."),
    ("🗄️", "Databricks MCP tools",
     "UC, Lakeflow, Genie, Lakebase and Apps provisioned from natural-language intent — "
     "infrastructure as conversation."),
    ("🧠", "Extraction prompt pattern",
     "One reusable ai_query() prompt turns every free-text MSL note into 9 governed columns — "
     "the product's core AI move."),
    ("📎", "Grounded-briefing guardrail",
     "The app's prompt enforces factual, context-only, non-promotional output — a safety "
     "pattern for regulated GenAI."),
    ("✅", "Generate → verify loop",
     "Every resource was validated live (execute_sql, endpoint checks) before moving on — "
     "AI proposes, we verify."),
    ("🧪", "Synthetic data generation",
     "Faker + Spark patterns produced realistic KOL, inquiry and engagement data at demo scale."),
]
for (icon_s, title_s, body_s), (icon, title, body) in zip(TRUST_CARDS, BUILT):
    set_text(builtai, icon_s, icon)
    set_text(builtai, title_s, title)
    set_text(builtai, body_s, body)
set_text(builtai, "Google Shape;463;p24", FOOTER)

# ---------------------------------------------------------------- 5. Reorder + renumber
# current prs.slides order (0-based):
#  0 title, 1 challenge, 2 built, 3 arch, 4 data, 5 ai, 6 genie, 7 dash,
#  8 copilot, 9 why-it-matters, 10 thankyou, 11 trust,
#  12 OUTCOME (clone), 13 TRADEOFFS (clone), 14 BUILT-WITH-AI (clone)
desired = [0, 12, 1, 2, 3, 4, 5, 6, 7, 8, 9, 13, 14, 11, 10]

sldIdLst = prs.slides._sldIdLst
sldIds = list(sldIdLst)
for sid in sldIds:
    sldIdLst.remove(sid)
for i in desired:
    sldIdLst.append(sldIds[i])

# renumber the top-right page badge (text box near left>=12in, top<0.6in)
RIGHT = Emu(int(12.0 * 914400))
TOPB = Emu(int(0.6 * 914400))
for pos, slide in enumerate(prs.slides, start=1):
    for sh in slide.shapes:
        if (sh.has_text_frame and sh.left is not None and sh.top is not None
                and sh.left >= RIGHT and sh.top <= TOPB):
            cur = sh.text_frame.text.strip()
            if cur.isdigit():
                set_text(slide, sh.name, f"{pos:02d}")
            break

prs.save(SRC)
print(f"saved {SRC} with {len(prs.slides)} slides")
for i, s in enumerate(prs.slides, 1):
    head = next((sh.text_frame.text.strip() for sh in s.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()), "")
    print(f"  {i:2d}. {head[:48]}")
