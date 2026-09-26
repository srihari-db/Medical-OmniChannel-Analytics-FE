"""Build the FE BAR deck by cloning the original Medical OmniChannel Analytics
PPTX (preserving its design system) and rewriting content for the Daiichi Sankyo
FE BAR build: rebrand to oncology, _sa701.moe -> _sa701.moa, add Lakeflow +
Lakebase, and inject the real numbers / evidence from this build.
"""
from pptx import Presentation

SRC = "/Users/srihari.a/Downloads/_Medical_OmniChannel_Analytics.pptx"
OUT = "/Users/srihari.a/febar/Medical-OmniChannel-Analytics-FE/deck/Medical_Omnichannel_Daiichi_Deck.pptx"

FOOTER = "Medical Omnichannel Intelligence  ·  Daiichi Sankyo (synthetic)  ·  Databricks  ·  _sa701.moa"
FOOTER_LONG = ("Medical Omnichannel Intelligence  ·  Daiichi Sankyo  ·  Databricks Data Intelligence Platform  "
               "·  adb-984752964297111 (Azure) · _sa701.moa  ·  synthetic demo data")

# (slide_index, shape_name) -> new text
EDITS = {
    # ---- Slide 1 : title ----
    (0, "Google Shape;87;p13"): "DATABRICKS  ·  DAIICHI SANKYO  ·  ONCOLOGY MEDICAL AFFAIRS",
    (0, "Google Shape;89;p13"): ("A unified, governed intelligence layer for Daiichi Sankyo oncology Medical Affairs — "
                                 "raw engagement signals flow through Lakeflow into one Unity Catalog 360° view, served "
                                 "operationally via Lakebase, made intelligent with Gen AI, queryable in Genie, and "
                                 "surfaced to MSLs in a Databricks App."),
    (0, "Google Shape;91;p13"): "Lakeflow → Unity Catalog",
    (0, "Google Shape;93;p13"): "Lakebase serving",
    (0, "Google Shape;95;p13"): "Genie + AI/BI",
    (0, "Google Shape;97;p13"): "Databricks App (Copilot)",

    # ---- Slide 2 : business challenge ----
    (1, "Google Shape;108;p14"): ("Daiichi Sankyo's MSLs engage the most influential oncology KOLs for Enhertu and "
                                  "Datroway — but the evidence of what a KOL actually cares about lives in disconnected "
                                  "systems, so follow-ups slip, medical inquiries age past SLA, and high-value KOLs go under-served."),
    (1, "Google Shape;125;p14"): FOOTER,

    # ---- Slide 3 : what we built ----
    (2, "Google Shape;132;p15"): "One governed platform — raw files to MSL action",
    (2, "Google Shape;135;p15"): ("Everything runs on the Databricks Data Intelligence Platform on one governed Unity "
                                  "Catalog schema (_sa701.moa). A Lakeflow medallion pipeline feeds Gold tables that power "
                                  "Genie, an AI/BI dashboard, a Lakebase-served MSL cockpit, and the Copilot app — no data "
                                  "copies, one security model."),
    (2, "Google Shape;141;p15"): ("Lakeflow medallion in Unity Catalog: 31 Delta tables across bronze, silver & gold, "
                                  "plus governed metric views and a Lakebase operational serving layer."),
    (2, "Google Shape;160;p15"): ("Shared foundation:  Lakeflow · Unity Catalog governance · Delta Lake · Lakebase "
                                  "(PostgreSQL) · Databricks-hosted LLM via ai_query() · serverless SQL"),
    (2, "Google Shape;161;p15"): FOOTER,

    # ---- Slide 4 : architecture ----
    (3, "Google Shape;176;p16"): "Lakeflow Auto Loader:  UC Volume → Bronze",
    (3, "Google Shape;207;p16"): ("🔒  Unity Catalog  —  single governance & lineage layer across every table, metric and "
                                  "surface  ·  Delta Lake  ·  Lakebase operational serving  ·  serverless SQL"),
    (3, "Google Shape;208;p16"): ("Genie, the AI/BI dashboard, the Lakebase MSL action queue, and the Copilot app all read "
                                  "the SAME governed Gold tables & metric views — one source of truth, one security model."),
    (3, "Google Shape;209;p16"): FOOTER,

    # ---- Slide 5 : data layer ----
    (4, "Google Shape;216;p17"): "Lakeflow medallion in Unity Catalog (_sa701.moa)",
    (4, "Google Shape;254;p17"): ("31 Delta tables  ·  2 governed metric views  ·  2 UC Volumes (raw + source_documents)  "
                                  "—  Lakeflow-built, Lakebase-served, one lineage graph"),
    (4, "Google Shape;255;p17"): FOOTER,

    # ---- Slide 6 : AI extraction (real example) ----
    (5, "Google Shape;268;p18"): ('"Met with the KOL to discuss real-world outcomes and treatment patterns post-T-DXd for '
                                  'Enhertu in HER2-low breast cancer. Tone positive, high scientific interest. Compared '
                                  'experience with Kadcyla and requested follow-up materials."'),
    (5, "Google Shape;275;p18"): "Real-world outcomes, T-DXd HER2-low",
    (5, "Google Shape;277;p18"): "Comparative real-world evidence",
    (5, "Google Shape;279;p18"): "Real-World Evidence",
    (5, "Google Shape;281;p18"): "Positive",
    (5, "Google Shape;283;p18"): "High",
    (5, "Google Shape;285;p18"): "true",
    (5, "Google Shape;287;p18"): "Kadcyla",
    (5, "Google Shape;289;p18"): "In-person",
    (5, "Google Shape;291;p18"): "Real-World Evidence",
    (5, "Google Shape;293;p18"): FOOTER,

    # ---- Slide 7 : Genie (verified questions) ----
    (6, "Google Shape;303;p19"): ('"Medical Omnichannel Intelligence — Daiichi Sankyo" Genie space lets medical & field '
                                  'teams self-serve governed answers — no SQL, no dashboard-hunting.'),
    (6, "Google Shape;307;p19"): "💬  Which 5 territories have the lowest 90-day HCP reach rate?",
    (6, "Google Shape;309;p19"): "💬  Which high-priority KOLs are under-engaged and overdue for an MSL visit?",
    (6, "Google Shape;311;p19"): "💬  Show medical inquiries overdue against SLA, by product.",
    (6, "Google Shape;313;p19"): "💬  What are the top scientific topics of interest across oncology KOLs?",
    (6, "Google Shape;315;p19"): "💬  List Tier 1 KOLs with open inquiries about Safety / ILD Management.",
    (6, "Google Shape;319;p19"): FOOTER,

    # ---- Slide 8 : dashboard ----
    (7, "Google Shape;329;p20"): ("A curated, always-on view of Medical Affairs engagement — 6 KPIs, reach by territory, "
                                  "topic interest, inquiry SLAs, priority tiers and the next-best-engagement queue — built "
                                  "on the same governed Gold tables and metric views that power Genie and the app."),
    (7, "Google Shape;355;p20"): FOOTER,

    # ---- Slide 9 : Copilot app ----
    (8, "Google Shape;365;p21"): ("Select a KOL and get a pre-engagement briefing that unites analytics, AI, governed "
                                  "content, and a next-best action — the MSL's single prep screen. The action queue is "
                                  "served from Lakebase; review decisions are written back to Lakebase."),
    (8, "Google Shape;370;p21"): "Prioritized queue served from Lakebase — highest MSL priority score first.",
    (8, "Google Shape;384;p21"): "Plan, approved content + Lakebase write-back",
    (8, "Google Shape;385;p21"): ("Objective, topic, channel, timing — plus an approved asset (with med review ID). The "
                                  "MSL logs the decision, written back transactionally to Lakebase. Requires MSL review."),
    (8, "Google Shape;386;p21"): ("Runs as a Databricks App (service-principal auth, no hardcoded tokens) reading governed "
                                  "Gold over serverless SQL and the Lakebase MSL action queue."),
    (8, "Google Shape;387;p21"): FOOTER,

    # ---- Slide 10 : why it matters (quantified) ----
    (9, "Google Shape;401;p22"): "AI briefings replace ~2 h/MSL/week of prep — est. ~$1.3M / year.",
    (9, "Google Shape;406;p22"): "Priority score + under-engaged flags lift priority-KOL reach — est. ~$1.5M / year.",
    (9, "Google Shape;411;p22"): "500+ free-text MSL notes become queryable topic-interest signal.",
    (9, "Google Shape;416;p22"): "Faster inquiry SLAs + approved-only content + full UC lineage — est. ~$0.6M / year.",
    (9, "Google Shape;419;p22"): ("One governed Lakehouse — Lakeflow to Lakebase to app — delivers analytics, natural-"
                                  "language self-service and an AI cockpit with no data copies and one security model. "
                                  "Estimated ~$3.4M / year for a 120-MSL field-medical org."),
    (9, "Google Shape;420;p22"): FOOTER,

    # ---- Slide 11 : thank you ----
    (10, "Google Shape;429;p23"): FOOTER_LONG,

    # ---- Slide 12 : trust ----
    (11, "Google Shape;450;p24"): ("One permission & lineage model across tables, metrics, Genie, dashboard, Lakebase and "
                                   "the app — auditable end to end; enforced UC tag policies, npi classified PII."),
    (11, "Google Shape;463;p24"): FOOTER,

    # ---- Slide 13 : demo flow ----
    (12, "Google Shape;483;p25"): "Pick a high-priority KOL → 360 → AI briefing → plan + approved content → log decision to Lakebase.",
    (12, "Google Shape;489;p25"): FOOTER_LONG,
}


def set_text(shape, new):
    tf = shape.text_frame
    first = tf.paragraphs[0]
    if first.runs:
        first.runs[0].text = new
        for r in first.runs[1:]:
            r.text = ""
    else:
        first.text = new
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""


prs = Presentation(SRC)
applied, missing = 0, []
for (si, name), new in EDITS.items():
    slide = prs.slides[si]
    shp = next((s for s in slide.shapes if s.name == name), None)
    if shp is None or not shp.has_text_frame:
        missing.append((si, name)); continue
    set_text(shp, new)
    applied += 1

prs.save(OUT)
print(f"applied {applied} edits; missing {missing}")
print("saved", OUT)
